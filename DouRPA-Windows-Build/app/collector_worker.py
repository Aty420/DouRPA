from __future__ import annotations

import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QThread, Signal

from app.collectors.douyin_mumu import DouyinMuMuCollector
from app.collectors.taobao import TaobaoCollector
from app.url_downloader import download_one, sanitize_filename


class ProductCollectorWorker(QThread):
    info = Signal(str)
    item_result = Signal(object)
    progress = Signal(int, int)
    finished_summary = Signal(object)

    def __init__(
        self,
        app_root: Path,
        keyword: str,
        platforms: dict[str, int],
        output_parent: Path,
        *,
        auto_download_images: bool = True,
        parent=None,
    ):
        super().__init__(parent)
        self.app_root = Path(app_root)
        self.keyword = str(keyword).strip()
        self.platforms = dict(platforms)
        self.output_parent = Path(output_parent)
        self.auto_download_images = bool(auto_download_images)
        self.stop_event = threading.Event()

    def request_stop(self):
        self.stop_event.set()

    def _emit_info(self, text: str):
        self.info.emit(text)

    def _save_excel(self, results: list[dict], path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = "商品采集结果"
        headers = ["商品标题", "首图URL", "SKU标题"]
        ws.append(headers)
        for result in results:
            sku = "；".join(result.get("sku_titles") or [])
            ws.append([
                result.get("title", ""),
                result.get("image_url", ""),
                sku,
            ])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B6FF5")
            cell.alignment = Alignment(horizontal="center")
        widths = [58, 72, 80]
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def run(self):
        safe_keyword = sanitize_filename(self.keyword, "商品")
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_dir = self.output_parent / f"{safe_keyword}_{stamp}"
        image_dir = output_dir / "首图"
        output_dir.mkdir(parents=True, exist_ok=True)
        image_dir.mkdir(parents=True, exist_ok=True)

        all_results: list[dict] = []
        errors: list[str] = []
        total_target = sum(max(0, int(v)) for v in self.platforms.values())

        for platform, count in self.platforms.items():
            if self.stop_event.is_set():
                break
            if count <= 0:
                continue
            try:
                if platform == "抖音":
                    collector = DouyinMuMuCollector(
                        progress_cb=self._emit_info,
                        stop_event=self.stop_event,
                    )
                elif platform == "淘宝":
                    collector = TaobaoCollector(
                        self.app_root / "data" / "collector_profiles" / "taobao",
                        progress_cb=self._emit_info,
                        stop_event=self.stop_event,
                    )
                else:
                    continue
                results = collector.collect(self.keyword, count)
            except Exception as exc:
                errors.append(f"{platform}: {exc}")
                self.info.emit(f"{platform}采集失败：{exc}")
                continue

            for result in results:
                if self.stop_event.is_set():
                    break
                all_results.append(result)
                status = "采集完成"
                if self.auto_download_images and result.get("image_url"):
                    try:
                        index = len(all_results)
                        download_one(
                            result["image_url"],
                            image_dir,
                            f"{index:03d}",
                            retries=2,
                            stop_event=self.stop_event,
                        )
                        status = "采集完成 · 首图已下载"
                    except Exception as exc:
                        status = f"采集完成 · 首图下载失败：{exc}"
                elif not result.get("image_url"):
                    status = "采集完成 · 未解析到首图URL"
                item = dict(result)
                item["status"] = status
                self.item_result.emit(item)
                self.progress.emit(len(all_results), total_target)

        excel_path = output_dir / "商品采集结果.xlsx"
        try:
            self._save_excel(all_results, excel_path)
        except Exception as exc:
            errors.append(f"Excel导出失败：{exc}")

        self.finished_summary.emit(
            {
                "output_dir": str(output_dir),
                "excel_path": str(excel_path),
                "count": len(all_results),
                "errors": errors,
                "stopped": self.stop_event.is_set(),
            }
        )
