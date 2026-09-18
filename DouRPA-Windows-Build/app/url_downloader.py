from __future__ import annotations

import html
import io
import json
import os
import re
import threading
import time
import urllib.parse
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image
from PySide6.QtCore import QThread, Signal


_INVALID_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def sanitize_filename(name: str, fallback: str = "image") -> str:
    name = _INVALID_FILENAME.sub("_", str(name or "")).strip(" ._")
    name = re.sub(r"\s+", " ", name)
    return (name[:120] or fallback).strip()


def _clean(value) -> str:
    value = html.unescape(str(value or ""))
    value = value.replace("\\u002F", "/").replace("\\u002f", "/")
    value = value.replace("\\u0026", "&").replace("\\/", "/")
    return re.sub(r"\s+", " ", value).strip(" \"'\\,")


def _profile_dir() -> Path:
    base = os.environ.get("LOCALAPPDATA")
    root = Path(base) / "DouRPA" if base else Path.home() / "AppData" / "Local" / "DouRPA"
    path = root / "data" / "url_parser_profile"
    path.mkdir(parents=True, exist_ok=True)
    return path


def _resolve_http_redirect(url: str) -> str:
    """Resolve the share short-link without opening a visible browser."""
    try:
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/138.0 Safari/537.36"
                ),
                "Accept-Language": "zh-CN,zh;q=0.9",
            },
        )
        with urllib.request.urlopen(req, timeout=15) as resp:
            return resp.geturl() or url
    except Exception:
        return url


def _extract_goods_detail_from_url(url: str) -> dict:
    """Extract title + first image from Douyin's goods_detail redirect parameter."""
    result = {
        "found": False,
        "title": "",
        "image_url": "",
        "image_urls": [],
    }

    try:
        parsed = urllib.parse.urlsplit(str(url or ""))
        query = urllib.parse.parse_qs(parsed.query, keep_blank_values=True)
        values = query.get("goods_detail") or []
        if not values:
            return result

        raw_value = values[0]
        obj = json.loads(raw_value) if isinstance(raw_value, str) else raw_value
        if not isinstance(obj, dict):
            return result

        result["found"] = True
        result["title"] = _clean(obj.get("title", ""))

        img = obj.get("img") or obj.get("image") or {}
        urls = []
        if isinstance(img, dict):
            candidate = img.get("url_list") or img.get("urls") or []
            if isinstance(candidate, str):
                candidate = [candidate]
            if isinstance(candidate, list):
                for item in candidate:
                    item = _clean(item)
                    if item.startswith(("http://", "https://")) and item not in urls:
                        urls.append(item)

        result["image_urls"] = urls
        result["image_url"] = urls[0] if urls else ""
    except Exception:
        pass

    return result


def _html_meta_value(source: str, *, property_name: str = "", name: str = "") -> str:
    if not source:
        return ""

    if property_name:
        patterns = [
            rf'<meta[^>]+property=["\']{re.escape(property_name)}["\'][^>]+content=["\']([^"\']+)["\']',
            rf'<meta[^>]+content=["\']([^"\']+)["\'][^>]+property=["\']{re.escape(property_name)}["\']',
        ]
    else:
        patterns = [
            rf'<meta[^>]+name=["\']{re.escape(name)}["\'][^>]+content=["\']([^"\']+)["\']',
            rf'<meta[^>]+content=["\']([^"\']+)["\'][^>]+name=["\']{re.escape(name)}["\']',
        ]

    for pattern in patterns:
        match = re.search(pattern, source, re.I | re.S)
        if match:
            return _clean(match.group(1))
    return ""


def _fetch_html_metadata(url: str) -> dict:
    """HTTP-only fallback for pages that expose OpenGraph metadata."""
    result = {"title": "", "image_url": ""}
    try:
        req = urllib.request.Request(
            url,
            headers={
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 Chrome/138.0 Safari/537.36"
                ),
                "Accept-Language": "zh-CN,zh;q=0.9",
            },
        )
        with urllib.request.urlopen(req, timeout=20) as resp:
            raw = resp.read(3_000_000)
            charset = resp.headers.get_content_charset() or "utf-8"
        source = raw.decode(charset, errors="ignore")

        result["title"] = (
            _html_meta_value(source, property_name="og:title")
            or _html_meta_value(source, name="twitter:title")
        )
        result["image_url"] = (
            _html_meta_value(source, property_name="og:image")
            or _html_meta_value(source, name="twitter:image")
        )

        if not result["title"]:
            match = re.search(r"<title[^>]*>(.*?)</title>", source, re.I | re.S)
            if match:
                result["title"] = _clean(re.sub(r"<[^>]+>", "", match.group(1)))
    except Exception:
        pass
    return result


def _browser_fallback(url: str) -> dict:
    """Headless fallback only when HTTP redirect/meta parsing cannot get title/image."""
    result = {"title": "", "image_url": ""}
    try:
        from playwright.sync_api import sync_playwright

        with sync_playwright() as p:
            context = p.chromium.launch_persistent_context(
                user_data_dir=str(_profile_dir()),
                channel="msedge",
                headless=True,
                viewport={"width": 1280, "height": 900},
                args=[
                    "--disable-blink-features=AutomationControlled",
                    "--no-first-run",
                    "--disable-notifications",
                ],
            )
            page = None
            try:
                page = context.new_page()
                page.goto(url, wait_until="domcontentloaded", timeout=45000)
                try:
                    page.wait_for_timeout(2500)
                except Exception:
                    pass

                # Title.
                try:
                    value = page.locator('meta[property="og:title"]').get_attribute("content")
                    result["title"] = _clean(value)
                except Exception:
                    pass
                if not result["title"]:
                    try:
                        result["title"] = _clean(page.title())
                    except Exception:
                        pass

                # First main-like large image.
                try:
                    result["image_url"] = page.evaluate(
                        """() => {
                          const bad = /(avatar|logo|icon|emoji|qrcode|qr-code|banner)/i;
                          const rows = [...document.images].map((img, idx) => ({
                            src: img.currentSrc || img.src || '',
                            w: img.naturalWidth || 0,
                            h: img.naturalHeight || 0,
                            area: (img.naturalWidth || 0) * (img.naturalHeight || 0),
                            idx
                          })).filter(x =>
                                /^https?:/i.test(x.src) &&
                                x.w >= 250 && x.h >= 250 &&
                                !bad.test(x.src))
                            .sort((a,b) => {
                              if (a.idx < 8 && b.idx >= 8) return -1;
                              if (b.idx < 8 && a.idx >= 8) return 1;
                              return b.area - a.area;
                            });
                          return rows.length ? rows[0].src : '';
                        }"""
                    ) or ""
                except Exception:
                    pass
            finally:
                try:
                    if page and not page.is_closed():
                        page.close()
                except Exception:
                    pass
                try:
                    context.close()
                except Exception:
                    pass
    except Exception:
        pass

    return result


def resolve_product_metadata(source_url: str, title_hint: str = "") -> dict:
    """Resolve only the fields the user still needs: title + first image URL."""
    resolved_url = _resolve_http_redirect(source_url)
    goods = _extract_goods_detail_from_url(resolved_url)

    title = _clean(title_hint) or _clean(goods.get("title", ""))
    image_url = _clean(goods.get("image_url", ""))

    # Douyin normally ends here: both fields come from the share text + redirect.
    if title and image_url:
        return {
            "title": title,
            "image_url": image_url,
            "final_url": resolved_url,
            "source": "share_text+goods_detail",
        }

    # HTTP metadata fallback.
    meta = _fetch_html_metadata(resolved_url)
    if not title:
        title = _clean(meta.get("title", ""))
    if not image_url:
        image_url = _clean(meta.get("image_url", ""))

    # Last resort: headless browser, still no SKU logic.
    if not title or not image_url:
        browser = _browser_fallback(resolved_url)
        if not title:
            title = _clean(browser.get("title", ""))
        if not image_url:
            image_url = _clean(browser.get("image_url", ""))

    return {
        "title": title,
        "image_url": image_url,
        "final_url": resolved_url,
        "source": "fallback",
    }


# Backward-compatible alias for any existing import/call site.
def resolve_page(context, source_url: str, index: int = 0, debug_dir: Path | None = None) -> dict:
    return resolve_product_metadata(source_url)


def _save_jpg_from_url(url: str, target: Path, retries=2):
    """Download source image and save a real JPEG."""
    last = None
    target = Path(target)

    for attempt in range(retries + 1):
        try:
            parsed = urllib.parse.urlparse(url)
            req = urllib.request.Request(
                url,
                headers={
                    "User-Agent": (
                        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 Chrome/138.0 Safari/537.36"
                    ),
                    "Referer": f"{parsed.scheme or 'https'}://{parsed.netloc}/",
                    "Accept": "image/avif,image/webp,image/apng,image/*,*/*;q=0.8",
                },
            )
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = resp.read()

            with Image.open(io.BytesIO(data)) as img:
                if getattr(img, "n_frames", 1) > 1:
                    img.seek(0)
                if img.mode != "RGB":
                    img = img.convert("RGB")

                target.parent.mkdir(parents=True, exist_ok=True)
                img.save(target, "JPEG", quality=95, optimize=True)

            if not target.exists() or target.stat().st_size <= 0:
                raise RuntimeError("JPG保存后文件为空")
            return

        except Exception as exc:
            last = exc
            if attempt < retries:
                time.sleep(0.7 * (attempt + 1))

    raise RuntimeError(f"JPG下载失败：{last}")


class ProductLinkWorker(QThread):
    item_result = Signal(int, object)
    info = Signal(str)
    progress = Signal(int, int)
    finished_summary = Signal(object)

    def __init__(
        self,
        urls: list[dict] | list[str],
        output_parent: Path,
        *,
        concurrency: int = 1,
        retries: int = 2,
        parent=None,
    ):
        super().__init__(parent)
        self.records = []
        for item in urls:
            if isinstance(item, dict):
                self.records.append({
                    "url": str(item.get("url") or "").strip(),
                    "title_hint": _clean(item.get("title_hint") or ""),
                })
            else:
                self.records.append({
                    "url": str(item or "").strip(),
                    "title_hint": "",
                })

        self.output_parent = Path(output_parent)
        self.concurrency = max(1, int(concurrency))
        self.retries = max(0, int(retries))
        self.stop_event = threading.Event()

    def request_stop(self):
        self.stop_event.set()

    @staticmethod
    def _save_excel(rows: list[dict], path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = "商品解析结果"
        ws.append(["商品标题", "首图URL"])

        for row in rows:
            ws.append([
                row.get("title", ""),
                row.get("image_url", ""),
            ])

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5B6FF5")
            cell.alignment = Alignment(horizontal="center")

        ws.column_dimensions["A"].width = 70
        ws.column_dimensions["B"].width = 100
        ws.freeze_panes = "A2"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)

    def run(self):
        stamp = time.strftime("%Y%m%d_%H%M%S")
        output_dir = self.output_parent / f"商品URL解析_{stamp}"
        image_dir = output_dir / "首图"
        output_dir.mkdir(parents=True, exist_ok=True)
        image_dir.mkdir(parents=True, exist_ok=True)

        rows = []
        errors = []
        total = len(self.records)

        self.info.emit("正在解析商品标题和首图，不再解析SKU。")

        for idx, record in enumerate(self.records):
            if self.stop_event.is_set():
                break

            url = record.get("url", "")
            title_hint = record.get("title_hint", "")
            self.info.emit(f"正在解析 {idx + 1}/{total}")

            try:
                meta = resolve_product_metadata(url, title_hint=title_hint)
                title = _clean(meta.get("title", ""))
                image_url = _clean(meta.get("image_url", ""))
                status = []

                if not title:
                    title = f"商品_{idx + 1:03d}"
                    status.append("未解析到商品标题")
                if not image_url:
                    status.append("未解析到首图URL")

                if image_url:
                    try:
                        filename = sanitize_filename(f"{idx + 1:03d}_{title}")
                        _save_jpg_from_url(
                            image_url,
                            image_dir / f"{filename}.jpg",
                            retries=self.retries,
                        )
                    except Exception as exc:
                        status.append(str(exc))

                row = {
                    "source_url": url,
                    "title": title,
                    "image_url": image_url,
                    "status": "完成" if not status else "；".join(status),
                }

            except Exception as exc:
                row = {
                    "source_url": url,
                    "title": title_hint or f"商品_{idx + 1:03d}",
                    "image_url": "",
                    "status": f"失败：{exc}",
                }
                errors.append(f"{idx + 1}: {exc}")

            rows.append(row)
            self.item_result.emit(idx, row)
            self.progress.emit(idx + 1, total)

        excel_path = output_dir / "商品解析结果.xlsx"
        try:
            self._save_excel(rows, excel_path)
        except Exception as exc:
            errors.append(f"Excel导出失败：{exc}")

        self.finished_summary.emit({
            "count": len(rows),
            "success": sum(1 for row in rows if row.get("image_url")),
            "failed": sum(1 for row in rows if not row.get("image_url")),
            "stopped": self.stop_event.is_set(),
            "errors": errors,
            "dest_dir": str(output_dir),
            "output_dir": str(output_dir),
            "excel_path": str(excel_path),
            "image_dir": str(image_dir),
        })
