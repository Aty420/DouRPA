from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import QSettings, Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QProgressBar,
    QPushButton,
    QSpinBox,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.main_window import button
from app.main_window_v220 import MainWindow as V220MainWindow
from app.url_downloader import ProductLinkWorker


class MainWindow(V220MainWindow):
    """V2.6.0: remove product-search collection and keep one URL product parser."""

    def __init__(self, root: Path):
        self.link_worker = None
        self.tool_settings = QSettings("LocalOps", "DouRPA")
        super().__init__(root)
        self.stack.addWidget(self._url_product_page())
        self.refresh_all()

    def _sidebar(self):
        panel = QFrame()
        panel.setObjectName("Sidebar")
        panel.setFixedWidth(270)

        lay = QVBoxLayout(panel)
        lay.setContentsMargins(20, 23, 18, 20)
        lay.setSpacing(8)

        brand_row = QHBoxLayout()
        brand_row.setSpacing(11)

        icon = QLabel()
        icon.setObjectName("BrandIcon")
        icon.setFixedSize(50, 50)
        icon.setAlignment(Qt.AlignCenter)
        icon.setPixmap(QApplication.windowIcon().pixmap(44, 44))

        brand_text = QVBoxLayout()
        brand_text.setSpacing(1)
        brand = QLabel("DouRPA Pro")
        brand.setObjectName("BrandName")
        sub = QLabel("让电商运营更简单")
        sub.setObjectName("BrandSub")
        brand_text.addWidget(brand)
        brand_text.addWidget(sub)

        brand_row.addWidget(icon)
        brand_row.addLayout(brand_text, 1)
        lay.addLayout(brand_row)
        lay.addSpacing(20)

        self.nav = []
        items = [
            ("⌂", "工作台", 0),
            ("▱", "裂变任务", 1),
            ("▢", "店铺管理", 2),
            ("◇", "源商品模板", 3),
            ("≡", "运行日志", 4),
            ("⚙", "系统设置", 5),
            ("⇩", "URL批量下载", 6),
        ]
        for glyph, text, idx in items:
            nav = QPushButton(f"  {glyph}     {text}")
            nav.setObjectName("NavButton")
            nav.setProperty("active", idx == 0)
            nav.setCursor(Qt.PointingHandCursor)
            nav.setFixedHeight(50)
            nav.clicked.connect(lambda _=False, i=idx: self.switch_page(i))
            lay.addWidget(nav)
            self.nav.append(nav)

        lay.addStretch(1)
        version = QLabel("DouRPA Pro  ·  v2.6.0")
        version.setObjectName("SidebarFooterPlain")
        version.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        lay.addWidget(version)
        return panel

    def _topbar(self):
        top = super()._topbar()
        self._titles = [
            ("工作台", "高效 · 稳定 · 智能  让重复的工作交给 DouRPA Pro"),
            ("裂变任务", "批量管理相似品发布任务"),
            ("店铺管理", "独立店铺 Profile 与登录状态"),
            ("源商品模板", "管理源商品定位与复用模板"),
            ("运行日志", "查看自动化节点、错误与恢复信息"),
            ("系统设置", "控制发布安全停点与页面配置"),
            ("URL批量下载", "上传商品链接 → 自动解析商品标题、首图URL、SKU标题 → 下载首图并导出Excel"),
        ]
        return top

    def _url_product_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(12)

        card = QFrame()
        card.setObjectName("Card")
        cl = QVBoxLayout(card)
        cl.setContentsMargins(18, 16, 18, 16)
        cl.setSpacing(10)

        head = QHBoxLayout()
        title_box = QVBoxLayout()
        title = QLabel("商品 URL 批量解析下载")
        title.setObjectName("SectionTitle")
        sub = QLabel(
            "无需商品搜索。把你自己复制的淘宝/抖音商品链接粘贴或放入 Excel，"
            "软件自动解析商品标题、首图URL、SKU标题，并下载第一张主图。"
        )
        sub.setObjectName("SectionSub")
        title_box.addWidget(title)
        title_box.addWidget(sub)
        head.addLayout(title_box)
        head.addStretch(1)

        paste = button("粘贴剪贴板")
        paste.clicked.connect(self._paste_links)
        import_txt = button("导入 TXT")
        import_txt.clicked.connect(self._import_txt)
        import_excel = button("上传 Excel", "primary")
        import_excel.clicked.connect(self._import_excel)
        head.addWidget(paste)
        head.addWidget(import_txt)
        head.addWidget(import_excel)
        cl.addLayout(head)

        self.link_input = QPlainTextEdit()
        self.link_input.setObjectName("UrlInputBox")
        self.link_input.setPlaceholderText(
            "每行一个商品链接，也可以直接粘贴抖音/淘宝分享文案，程序会自动提取其中的 http/https 链接。\n\n"
            "Excel 推荐列名：商品链接 / URL / 链接 / 分享链接"
        )
        self.link_input.setMinimumHeight(150)
        cl.addWidget(self.link_input)

        settings = QHBoxLayout()
        default_dir = self.tool_settings.value(
            "product_url_output_dir", str(self.root / "URL下载")
        )
        self.output_dir = QLineEdit(str(default_dir))
        choose = button("选择目录")
        choose.clicked.connect(self._choose_output_dir)

        self.concurrency = QSpinBox()
        self.concurrency.setRange(1, 6)
        self.concurrency.setValue(
            self.tool_settings.value("product_url_concurrency", 3, type=int)
        )
        self.retries = QSpinBox()
        self.retries.setRange(0, 5)
        self.retries.setValue(
            self.tool_settings.value("product_url_retries", 2, type=int)
        )

        settings.addWidget(QLabel("保存目录"))
        settings.addWidget(self.output_dir, 1)
        settings.addWidget(choose)
        settings.addSpacing(12)
        settings.addWidget(QLabel("并发"))
        settings.addWidget(self.concurrency)
        settings.addWidget(QLabel("失败重试"))
        settings.addWidget(self.retries)
        cl.addLayout(settings)

        actions = QHBoxLayout()
        self.start_btn = button("开始解析并下载", "primary")
        self.start_btn.clicked.connect(self._start)
        self.stop_btn = button("停止", "danger")
        self.stop_btn.setEnabled(False)
        self.stop_btn.clicked.connect(self._stop)
        self.info_label = QLabel("等待粘贴或上传商品链接")
        self.info_label.setObjectName("SectionSub")
        self.progress = QProgressBar()
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setFixedWidth(220)

        actions.addWidget(self.start_btn)
        actions.addWidget(self.stop_btn)
        actions.addSpacing(12)
        actions.addWidget(self.info_label, 1)
        actions.addWidget(self.progress)
        cl.addLayout(actions)

        lay.addWidget(card)

        result = QFrame()
        result.setObjectName("Card")
        rl = QVBoxLayout(result)
        rl.setContentsMargins(16, 14, 16, 14)

        rt = QLabel("实时解析结果")
        rt.setObjectName("SectionTitle")
        rl.addWidget(rt)

        self.result_table = self._table(
            ["商品链接", "商品标题", "首图URL", "SKU标题", "状态"]
        )
        self.result_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.result_table.setColumnWidth(0, 340)
        self.result_table.setColumnWidth(1, 360)
        self.result_table.setColumnWidth(2, 420)
        self.result_table.setColumnWidth(3, 430)
        self.result_table.setColumnWidth(4, 240)
        rl.addWidget(self.result_table)
        lay.addWidget(result, 1)
        return page

    def _paste_links(self):
        text = QApplication.clipboard().text().strip()
        if not text:
            return
        current = self.link_input.toPlainText().strip()
        self.link_input.setPlainText((current + "\n" + text).strip() if current else text)

    def _import_txt(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "导入商品链接 TXT", "", "Text (*.txt);;All (*.*)"
        )
        if not path:
            return
        self.link_input.setPlainText(
            Path(path).read_text(encoding="utf-8", errors="ignore")
        )

    def _import_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "上传商品链接 Excel", "", "Excel (*.xlsx *.xlsm)"
        )
        if not path:
            return

        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise RuntimeError("Excel 为空")

            headers = [str(x or "").strip() for x in rows[0]]
            preferred = (
                "商品链接", "商品URL", "产品链接", "分享链接",
                "URL", "url", "链接",
            )

            url_idx = -1
            for name in preferred:
                for i, header in enumerate(headers):
                    if header == name or name.lower() == header.lower():
                        url_idx = i
                        break
                if url_idx >= 0:
                    break

            links = []
            data_rows = rows[1:] if url_idx >= 0 else rows
            for row in data_rows:
                values = list(row)
                candidates = []
                if 0 <= url_idx < len(values):
                    candidates.append(str(values[url_idx] or ""))
                else:
                    candidates.extend(str(v or "") for v in values)

                for value in candidates:
                    match = re.search(r"https?://[^\s，。；;]+", value)
                    if match:
                        links.append(match.group(0).rstrip(")】]"))
                        break

            self.link_input.setPlainText("\n".join(links))
            self.info_label.setText(f"已从 Excel 导入 {len(links)} 个商品链接")
        except Exception as exc:
            QMessageBox.critical(self, "Excel 导入失败", str(exc))

    def _choose_output_dir(self):
        path = QFileDialog.getExistingDirectory(
            self, "选择保存目录", self.output_dir.text()
        )
        if path:
            self.output_dir.setText(path)
            self.tool_settings.setValue("product_url_output_dir", path)

    def _parse_links(self) -> list[str]:
        text = self.link_input.toPlainText()
        links = re.findall(r"https?://[^\s，。；;]+", text)
        result = []
        seen = set()
        for link in links:
            link = link.rstrip(")】]}>\"'")
            if link and link not in seen:
                seen.add(link)
                result.append(link)
        return result

    def _start(self):
        links = self._parse_links()
        if not links:
            QMessageBox.information(
                self,
                "没有商品链接",
                "请粘贴商品链接，或上传包含“商品链接 / URL / 链接”列的 Excel。",
            )
            return
        if self.link_worker and self.link_worker.isRunning():
            return

        output = Path(
            self.output_dir.text().strip() or self.root / "URL下载"
        )
        output.mkdir(parents=True, exist_ok=True)
        self.tool_settings.setValue("product_url_output_dir", str(output))
        self.tool_settings.setValue("product_url_concurrency", self.concurrency.value())
        self.tool_settings.setValue("product_url_retries", self.retries.value())

        self.result_table.setRowCount(len(links))
        for row, link in enumerate(links):
            self.result_table.setItem(row, 0, QTableWidgetItem(link))
            self.result_table.setItem(row, 1, QTableWidgetItem(""))
            self.result_table.setItem(row, 2, QTableWidgetItem(""))
            self.result_table.setItem(row, 3, QTableWidgetItem(""))
            self.result_table.setItem(row, 4, QTableWidgetItem("等待"))

        self.progress.setValue(0)
        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.info_label.setText(f"正在处理 {len(links)} 个商品链接…")

        worker = ProductLinkWorker(
            links,
            output,
            concurrency=self.concurrency.value(),
            retries=self.retries.value(),
            parent=self,
        )
        self.link_worker = worker
        worker.item_result.connect(self._item_result)
        worker.info.connect(self.info_label.setText)
        worker.progress.connect(self._progress)
        worker.finished_summary.connect(self._finished)
        worker.start()

    def _stop(self):
        if self.link_worker and self.link_worker.isRunning():
            self.link_worker.request_stop()
            self.stop_btn.setEnabled(False)
            self.info_label.setText("正在停止…")

    def _item_result(self, index: int, item: dict):
        if index < 0 or index >= self.result_table.rowCount():
            return
        self.result_table.setItem(index, 0, QTableWidgetItem(item.get("source_url", "")))
        self.result_table.setItem(index, 1, QTableWidgetItem(item.get("title", "")))
        self.result_table.setItem(index, 2, QTableWidgetItem(item.get("image_url", "")))
        self.result_table.setItem(
            index,
            3,
            QTableWidgetItem("；".join(item.get("sku_titles") or [])),
        )
        self.result_table.setItem(index, 4, QTableWidgetItem(item.get("status", "")))

    def _progress(self, done: int, total: int):
        value = int(done * 100 / total) if total else 0
        self.progress.setValue(max(0, min(100, value)))

    def _finished(self, summary: dict):
        self.start_btn.setEnabled(True)
        self.stop_btn.setEnabled(False)
        if not summary.get("stopped"):
            self.progress.setValue(100)

        self.info_label.setText(
            f"完成 {summary.get('count', 0)} 条 · Excel：{summary.get('excel_path', '')}"
        )

        text = (
            f"处理完成：{summary.get('count', 0)} 个商品链接\n"
            f"解析到首图：{summary.get('success', 0)} 个\n"
            f"未解析到首图：{summary.get('failed', 0)} 个\n\n"
            f"Excel：{summary.get('excel_path', '')}\n"
            f"首图文件夹：{summary.get('image_dir', '')}"
        )
        if summary.get("errors"):
            text += "\n\n异常：\n" + "\n".join(summary["errors"][:8])

        QMessageBox.information(
            self,
            "处理完成" if not summary.get("stopped") else "处理已停止",
            text,
        )
