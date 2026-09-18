from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QPalette
from PySide6.QtWidgets import (
    QAbstractButton,
    QComboBox,
    QDoubleSpinBox,
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QScrollArea,
    QSpinBox,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)

from app.main_window import SectionCard, button
from app.main_window_v240 import MainWindow as V240MainWindow
from app.url_downloader import ProductLinkWorker
from app.ui_tokens import SIZE


_SUPPORTED_PRODUCT_HOSTS = (
    "v.douyin.com",
    "www.douyin.com",
    "douyin.com",
    "haohuo.jinritemai.com",
    "jinritemai.com",
    "m.tb.cn",
    "item.taobao.com",
    "taobao.com",
    "detail.tmall.com",
    "tmall.com",
)

_BLOCKED_ASSET_SUFFIXES = (
    ".zip", ".js", ".css", ".json", ".wasm", ".bin",
    ".apk", ".exe", ".dll", ".woff", ".woff2", ".ttf",
)


def _is_supported_product_url(url: str) -> bool:
    from urllib.parse import urlparse

    try:
        parsed = urlparse(str(url or "").strip())
        host = (parsed.hostname or "").lower()
        path = (parsed.path or "").lower()
    except Exception:
        return False

    if not host:
        return False
    if path.endswith(_BLOCKED_ASSET_SUFFIXES):
        return False
    if any(bad in host for bad in (
        "static", "cdn", "verifycenter", "rc-verify", "byteimg", "snssdk"
    )):
        return False

    return any(
        host == allowed or host.endswith("." + allowed)
        for allowed in _SUPPORTED_PRODUCT_HOSTS
    )


def _url_score(url: str) -> int:
    from urllib.parse import urlparse

    try:
        parsed = urlparse(url)
        host = (parsed.hostname or "").lower()
        path = (parsed.path or "").lower()
    except Exception:
        return -999

    score = 0
    if host == "v.douyin.com":
        score += 100
    elif host == "m.tb.cn":
        score += 95
    elif "item.taobao.com" in host:
        score += 90
    elif "detail.tmall.com" in host:
        score += 90
    elif "douyin.com" in host:
        score += 80
    elif "jinritemai.com" in host:
        score += 75

    if any(k in path for k in ("/item", "/product", "/detail", "/goods")):
        score += 20
    return score


def _extract_records(text: str) -> list[dict]:
    text = str(text or "")
    matches = list(re.finditer(r"https?://[^\s，。；;]+", text))

    candidates = []
    for match in matches:
        url = match.group(0).rstrip(")】]}>\"'")
        if _is_supported_product_url(url):
            candidates.append((match, url))

    result = []
    seen = set()

    for i, (match, url) in enumerate(candidates):
        if url in seen:
            continue

        end = candidates[i + 1][0].start() if i + 1 < len(candidates) else len(text)
        tail = text[match.end():end]
        tail = re.split(
            r"(?:\r?\n|长按复制此条消息|打开抖音搜索|查看商品详情)",
            tail,
            maxsplit=1,
        )[0]
        title_hint = tail.strip(" \t\r\n：:，,。；;")
        title_hint = re.sub(
            r"^(?:【抖音商城】|【淘宝】|【天猫】)\s*",
            "",
            title_hint,
        ).strip()

        if not title_hint or title_hint.startswith(("http://", "https://")):
            title_hint = ""

        seen.add(url)
        result.append({
            "url": url,
            "title_hint": title_hint,
            "_score": _url_score(url),
        })

    result.sort(key=lambda x: x.get("_score", 0), reverse=True)
    for item in result:
        item.pop("_score", None)
    return result


class MainWindow(V240MainWindow):
    """V2.7.2: readable typography/spacing + title/image URL parser + clear logs."""

    def __init__(self, root: Path):
        super().__init__(root)

        # 1280×720 must be allowed. On short screens the workbench scrolls
        # vertically instead of compressing/cropping fixed-height cards.
        self.setMinimumSize(1180, 680)
        self._install_dashboard_scroll()
        self._apply_readability_policy()
        self._apply_responsive_spacing()

    def _install_dashboard_scroll(self):
        if not hasattr(self, "stack") or self.stack.count() == 0:
            return

        original = self.stack.widget(0)
        if isinstance(original, QScrollArea):
            self._dashboard_scroll = original
            return

        current_index = self.stack.currentIndex()
        self.stack.removeWidget(original)

        scroll = QScrollArea()
        scroll.setObjectName("DashboardScroll")
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)

        # Keeps the reference composition intact on 1920×1080, while a
        # 1280×720 window simply gets a vertical scrollbar instead of clipping.
        original.setMinimumHeight(790)
        scroll.setWidget(original)
        self.stack.insertWidget(0, scroll)
        self._dashboard_scroll = scroll

        self.stack.setCurrentIndex(max(0, current_index))

    def _apply_readability_policy(self):
        # Tables: minimum 44px rows everywhere, including dashboard/task/store/URL.
        for table in self.findChildren(QTableWidget):
            table.verticalHeader().setMinimumSectionSize(SIZE["table_row_height"])
            table.verticalHeader().setDefaultSectionSize(SIZE["table_row_height"])
            table.horizontalHeader().setMinimumHeight(SIZE["table_header_height"])

        # Input/placeholder/disabled contrast.
        for cls in (QLineEdit, QComboBox, QSpinBox, QDoubleSpinBox, QPlainTextEdit):
            for widget in self.findChildren(cls):
                palette = widget.palette()
                palette.setColor(QPalette.PlaceholderText, QColor("#7A869C"))
                palette.setColor(QPalette.Disabled, QPalette.Text, QColor("#929CAF"))
                palette.setColor(QPalette.Disabled, QPalette.WindowText, QColor("#929CAF"))
                widget.setPalette(palette)

        # Standard action buttons should never become text-height-only controls.
        for btn in self.findChildren(QAbstractButton):
            if btn.objectName() in {"PrimaryButton", "SecondaryButton", "DangerButton"}:
                btn.setMinimumHeight(SIZE["control_height"])

    def _apply_responsive_spacing(self):
        if not self.centralWidget() or not self.centralWidget().layout():
            return

        outer = self.centralWidget().layout()
        sidebar = outer.itemAt(0).widget() if outer.count() > 0 else None
        content = outer.itemAt(1).widget() if outer.count() > 1 else None

        compact = self.width() < 1400

        if sidebar is not None:
            sidebar.setFixedWidth(
                SIZE["compact_sidebar_width"]
                if compact
                else SIZE["normal_sidebar_width"]
            )

        if content is not None and content.layout() is not None:
            if compact:
                content.layout().setContentsMargins(18, 14, 18, 16)
                content.layout().setSpacing(12)
            else:
                content.layout().setContentsMargins(28, 18, 28, 22)
                content.layout().setSpacing(14)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply_responsive_spacing()

    def _topbar(self):
        top = super()._topbar()
        if len(getattr(self, "_titles", [])) >= 7:
            self._titles[6] = (
                "URL批量下载",
                "上传商品链接 → 自动解析商品标题与首图URL → 下载首图并导出Excel",
            )
        return top

    def _logs_page(self):
        """Run-log page with refresh + clear actions."""
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setContentsMargins(0, 0, 0, 0)

        card = SectionCard(
            "运行日志",
            "RPA 页面结构变化时，优先查看失败节点、截图和 config/selectors.json。",
        )

        self.log_edit = QPlainTextEdit()
        self.log_edit.setReadOnly(True)
        card.body.addWidget(self.log_edit)

        actions = QHBoxLayout()
        refresh_btn = button("刷新日志")
        refresh_btn.clicked.connect(self.refresh_logs)

        clear_btn = button("清除日志", "danger")
        clear_btn.clicked.connect(self._clear_run_logs)

        actions.addWidget(refresh_btn)
        actions.addWidget(clear_btn)
        actions.addStretch(1)
        card.body.addLayout(actions)

        lay.addWidget(card, 1)
        return page

    def _clear_run_logs(self):
        """Delete only rows in the logs table; keep tasks/stores/templates intact."""
        answer = QMessageBox.question(
            self,
            "清除运行日志",
            "确认清除全部运行日志？\n\n"
            "此操作只清除“运行日志”记录，不会删除裂变任务、店铺、源商品模板、"
            "浏览器登录状态或截图文件。",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        )
        if answer != QMessageBox.Yes:
            return

        try:
            with self.db.connect() as conn:
                conn.execute("DELETE FROM logs")

            self.refresh_logs()
            QMessageBox.information(self, "清除完成", "运行日志已清空。")
        except Exception as exc:
            QMessageBox.critical(self, "清除失败", f"无法清除运行日志：{exc}")

    def _url_product_page(self):
        page = super()._url_product_page()

        # Remove SKU wording from existing explanatory labels.
        for label in page.findChildren(QLabel):
            value = label.text()
            if "SKU标题" in value:
                value = value.replace("、SKU标题", "").replace("SKU标题、", "")
                value = value.replace("SKU标题", "")
                value = re.sub(r"、{2,}", "、", value)
                label.setText(value)

        # Result table is now: link / title / first-image URL / status.
        self.result_table.setColumnCount(4)
        self.result_table.setHorizontalHeaderLabels(
            ["商品链接", "商品标题", "首图URL", "状态"]
        )
        self.result_table.setColumnWidth(0, 360)
        self.result_table.setColumnWidth(1, 420)
        self.result_table.setColumnWidth(2, 520)
        self.result_table.setColumnWidth(3, 240)
        return page

    def _import_excel(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "上传商品链接 Excel", "", "Excel (*.xlsx *.xlsm)"
        )
        if not path:
            return

        try:
            wb = load_workbook(path, data_only=True, read_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise RuntimeError("Excel 为空")

            headers = [str(x or "").strip() for x in rows[0]]
            preferred = (
                "商品链接", "商品URL", "产品链接", "分享链接", "分享文案",
                "URL", "url", "链接",
            )

            url_idx = -1
            for name in preferred:
                for i, header in enumerate(headers):
                    if header == name or name.lower() == header.lower():
                        url_idx = i
                        break
                if url_idx >= 0:
                    break

            values_with_links = []
            data_rows = rows[1:] if url_idx >= 0 else rows

            for row in data_rows:
                values = list(row)
                candidates = []

                if 0 <= url_idx < len(values):
                    candidates.append(str(values[url_idx] or ""))
                else:
                    candidates.extend(str(v or "") for v in values)

                for value in candidates:
                    if _extract_records(value):
                        values_with_links.append(value.strip())
                        break

            self.link_input.setPlainText("\n\n".join(values_with_links))
            records = _extract_records(self.link_input.toPlainText())
            preview = records[0]["url"] if records else "未识别到有效商品链接"
            self.info_label.setText(
                f"已从 Excel 导入 {len(records)} 个商品链接 · 首条：{preview}"
            )

        except Exception as exc:
            QMessageBox.critical(self, "Excel 导入失败", str(exc))

    def _parse_links(self) -> list[dict]:
        return _extract_records(self.link_input.toPlainText())

    def _start(self):
        records = self._parse_links()

        if not records:
            QMessageBox.information(
                self,
                "没有商品链接",
                "请粘贴商品分享文案/链接，或上传包含“商品链接 / 分享文案 / URL / 链接”列的 Excel。",
            )
            return

        if self.link_worker and self.link_worker.isRunning():
            return

        output = Path(
            self.output_dir.text().strip() or self.root / "URL下载"
        )
        output.mkdir(parents=True, exist_ok=True)

        self.tool_settings.setValue("product_url_output_dir", str(output))
        self.tool_settings.setValue("product_url_retries", self.retries.value())

        self.result_table.setRowCount(len(records))

        for row, record in enumerate(records):
            self.result_table.setItem(row, 0, QTableWidgetItem(record.get("url", "")))
            self.result_table.setItem(
                row, 1, QTableWidgetItem(record.get("title_hint", ""))
            )
            self.result_table.setItem(row, 2, QTableWidgetItem(""))
            self.result_table.setItem(row, 3, QTableWidgetItem("等待"))

        self.progress.setValue(0)
        self.start_btn.setEnabled(False)
        self.stop_btn.setEnabled(True)
        self.info_label.setText(
            f"正在处理 {len(records)} 个商品链接；仅解析商品标题和首图"
        )

        worker = ProductLinkWorker(
            records,
            output,
            concurrency=self.concurrency.value(),
            retries=self.retries.value(),
            parent=self,
        )
        self.link_worker = worker
        worker.item_result.connect(self._item_result)
        worker.info.connect(self.info_label.setText)
        worker.progress.connect(self._progress)
        worker.finished_summary.connect(self._finished)
        worker.start()

    def _item_result(self, index: int, item: dict):
        if index < 0 or index >= self.result_table.rowCount():
            return

        self.result_table.setItem(
            index, 0, QTableWidgetItem(item.get("source_url", ""))
        )
        self.result_table.setItem(
            index, 1, QTableWidgetItem(item.get("title", ""))
        )
        self.result_table.setItem(
            index, 2, QTableWidgetItem(item.get("image_url", ""))
        )
        self.result_table.setItem(
            index, 3, QTableWidgetItem(item.get("status", ""))
        )
